"""ASYS - Apartman ve Site Yönetim Sistemi V1 Backend."""
from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

import os
import io
import uuid
import bcrypt
import jwt
import logging
import re
import ipaddress
from datetime import datetime, timezone, timedelta, date
from typing import List, Optional, Literal, Any, Dict
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse

from fastapi import FastAPI, APIRouter, Depends, HTTPException, Request, Response
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr, ConfigDict
import httpx

# ---------- Setup ----------
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]
JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALG = "HS256"
ACCESS_TTL_MIN = 60 * 12  # 12 hours

EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY", "")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "ASYS")
FRONTEND_URL = os.environ.get("FRONTEND_URL", "").rstrip("/")

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

app = FastAPI(title="ASYS API")
api = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("asys")


# ---------- Helpers ----------
def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def new_id() -> str:
    return str(uuid.uuid4())


def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()


def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False


def create_token(user_id: str, role: str) -> str:
    payload = {
        "sub": user_id,
        "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(minutes=ACCESS_TTL_MIN),
        "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)


async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Kimlik doğrulama gerekli")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Oturum süresi doldu")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Geçersiz token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user or not user.get("active", True):
        raise HTTPException(401, "Kullanıcı bulunamadı")
    return user


def require_role(*roles):
    async def _dep(user: dict = Depends(get_current_user)):
        if user["role"] not in roles:
            raise HTTPException(403, "Bu işlem için yetkiniz yok")
        return user
    return _dep


async def audit(user: dict, action: str, entity: str, entity_id: str = "", meta: Optional[dict] = None, site_id: str = ""):
    await db.audit_logs.insert_one({
        "id": new_id(), "ts": now_iso(),
        "user_id": user["id"], "user_email": user["email"], "user_role": user["role"],
        "action": action, "entity": entity, "entity_id": entity_id,
        "site_id": site_id, "meta": meta or {},
    })


async def check_period_open(site_id: str, target_date: str):
    """Ensure the period for target_date is open. target_date is ISO date str."""
    d = target_date[:10]
    year, month = int(d[:4]), int(d[5:7])
    period = await db.periods.find_one({"site_id": site_id, "year": year, "month": month})
    if period and period.get("status") == "closed":
        raise HTTPException(400, f"{year}-{month:02d} dönemi kapalı. Bu tarihe kayıt yapılamaz.")


# ---------- Models ----------
class SetupAdmin(BaseModel):
    email: EmailStr
    name: str
    password: str = Field(min_length=6)
    phone: Optional[str] = ""

class LoginIn(BaseModel):
    email: EmailStr
    password: str

class RegisterUserIn(BaseModel):
    email: EmailStr
    name: str
    password: str = Field(min_length=6)
    role: Literal["admin", "muhasebe", "denetci"]
    phone: Optional[str] = ""

class UpdateUserIn(BaseModel):
    name: Optional[str] = None
    role: Optional[Literal["admin", "muhasebe", "denetci"]] = None
    active: Optional[bool] = None
    password: Optional[str] = None
    phone: Optional[str] = None

class SiteIn(BaseModel):
    name: str
    address: Optional[str] = ""
    tax_no: Optional[str] = ""

class BlockIn(BaseModel):
    site_id: str
    name: str

class UnitIn(BaseModel):
    site_id: str
    block_id: str
    no: str
    kind: Literal["konut", "isyeri"] = "konut"
    area_m2: Optional[float] = 0
    share_ratio: Optional[float] = 0
    active: bool = True


class UnitPatchIn(BaseModel):
    block_id: Optional[str] = None
    no: Optional[str] = None
    kind: Optional[Literal["konut", "isyeri"]] = None
    area_m2: Optional[float] = None
    share_ratio: Optional[float] = None
    active: Optional[bool] = None

class PersonIn(BaseModel):
    site_id: str
    kind: Literal["malik", "kiraci"]
    name: str
    phone: Optional[str] = ""
    email: Optional[str] = ""
    tc_no: Optional[str] = ""
    unit_id: Optional[str] = ""
    active: bool = True


class PersonPatchIn(BaseModel):
    kind: Optional[Literal["malik", "kiraci"]] = None
    name: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    tc_no: Optional[str] = None
    unit_id: Optional[str] = None
    active: Optional[bool] = None

class RelationIn(BaseModel):
    site_id: str
    unit_id: str
    person_id: str
    kind: Literal["malik", "kiraci"]
    start_date: str
    end_date: Optional[str] = None

class AccrualBatchIn(BaseModel):
    site_id: str
    period: str  # YYYY-MM
    due_date: str  # ISO
    description: str = "Aidat"
    method: Literal["esit", "metrekare", "arsa_payi"] = "esit"
    amount_total_kurus: int  # for esit: per unit; for others: total distributed

class AccrualExtraIn(BaseModel):
    site_id: str
    unit_id: str
    due_date: str
    description: str
    amount_kurus: int

class ReverseAccrualIn(BaseModel):
    reason: str

class CollectionIn(BaseModel):
    site_id: str
    unit_id: str
    account_id: str  # cash or bank account id
    account_kind: Literal["cash", "bank"]
    date: str
    amount_kurus: int
    reference: Optional[str] = ""  # bank ref
    note: Optional[str] = ""

class ExpenseIn(BaseModel):
    site_id: str
    account_id: str
    account_kind: Literal["cash", "bank"]
    date: str
    amount_kurus: int
    category: str
    vendor: Optional[str] = ""
    invoice_no: Optional[str] = ""
    description: Optional[str] = ""
    reference: Optional[str] = ""

class AccountIn(BaseModel):
    site_id: str
    kind: Literal["cash", "bank"]
    name: str
    iban: Optional[str] = ""
    bank_name: Optional[str] = ""
    opening_balance_kurus: int = 0

class TransferIn(BaseModel):
    site_id: str
    from_id: str
    from_kind: Literal["cash", "bank"]
    to_id: str
    to_kind: Literal["cash", "bank"]
    date: str
    amount_kurus: int
    reference: Optional[str] = ""
    note: Optional[str] = ""

class PeriodActionIn(BaseModel):
    site_id: str
    year: int
    month: int
    reason: Optional[str] = ""


class ChangePasswordIn(BaseModel):
    current_password: str
    new_password: str = Field(min_length=6)


class ForgotPasswordIn(BaseModel):
    email: EmailStr


class ResetPasswordIn(BaseModel):
    token: str
    new_password: str = Field(min_length=6)


# ---------- Email guardrail gate (from Resend playbook) ----------
_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "reply with the code", "send your password", "cvv",
             "send us your password", "enter your password below", "confirm your card number",
             "your full card number", "seed phrase", "recovery phrase", "verify your card",
             "social security number", "confirm your bank details")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)


def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)


def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)


class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []
    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []


def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan(); scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks the recipient for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links/assets must be absolute https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Shortened, numeric-host or credential-bearing URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text {m.group(1)!r} != real link host {real!r} (G3)")


async def send_email(*, to: str, subject: str, html: str) -> Optional[str]:
    if not EMAIL_KEY:
        log.warning("EMERGENT_EMAIL_KEY not set - skipping email send")
        return None
    _assert_safe_email(subject, html)
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    try:
        async with httpx.AsyncClient(timeout=30) as c:
            resp = await c.post(
                f"{EMAIL_BASE_URL}/api/v1/email/send",
                headers={"X-Email-Key": EMAIL_KEY},
                json=payload,
            )
        resp.raise_for_status()
        return resp.json().get("id")
    except httpx.HTTPStatusError as e:
        log.error(f"Email send failed: {e.response.status_code} {e.response.text}")
        raise HTTPException(status_code=502, detail="E-posta gönderilemedi")
    except Exception as e:
        log.error(f"Email send error: {e}")
        raise HTTPException(status_code=500, detail="E-posta gönderilemedi")


# ---------- Startup ----------
@app.on_event("startup")
async def _startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("id", unique=True)
    await db.sites.create_index("id", unique=True)
    await db.blocks.create_index([("site_id", 1), ("id", 1)])
    await db.units.create_index([("site_id", 1), ("id", 1)])
    await db.persons.create_index([("site_id", 1), ("id", 1)])
    await db.accruals.create_index([("site_id", 1), ("unit_id", 1)])
    await db.payments.create_index([("site_id", 1), ("unit_id", 1)])
    # unique bank reference per site+account
    await db.account_transactions.create_index(
        [("site_id", 1), ("account_id", 1), ("account_kind", 1), ("reference", 1)],
        unique=True,
        partialFilterExpression={"reference": {"$gt": ""}},
    )
    await db.journal_entries.create_index([("site_id", 1), ("date", 1)])
    await db.audit_logs.create_index([("ts", -1)])
    await db.periods.create_index([("site_id", 1), ("year", 1), ("month", 1)], unique=True)
    log.info("ASYS başlatıldı")


@app.on_event("shutdown")
async def _shutdown():
    client.close()


# ---------- Setup / Auth ----------
@api.get("/setup/status")
async def setup_status():
    admin_count = await db.users.count_documents({"role": "admin"})
    return {"setup_required": admin_count == 0}


@api.post("/setup/admin")
async def setup_admin(inp: SetupAdmin, response: Response):
    if await db.users.count_documents({"role": "admin"}) > 0:
        raise HTTPException(400, "Kurulum zaten tamamlanmış")
    user = {
        "id": new_id(),
        "email": inp.email.lower(),
        "name": inp.name,
        "role": "admin",
        "active": True,
        "phone": inp.phone or "",
        "password_hash": hash_pw(inp.password),
        "created_at": now_iso(),
    }
    try:
        await db.users.insert_one(user)
    except Exception:
        raise HTTPException(400, "Bu e-posta zaten kullanımda")
    user.pop("_id", None)
    token = create_token(user["id"], "admin")
    response.set_cookie("access_token", token, httponly=True, samesite="lax", max_age=ACCESS_TTL_MIN * 60, path="/")
    return {"user": {k: v for k, v in user.items() if k != "password_hash"}, "access_token": token}


@api.post("/auth/login")
async def login(inp: LoginIn, response: Response):
    user = await db.users.find_one({"email": inp.email.lower()})
    if not user or not user.get("active", True) or not verify_pw(inp.password, user["password_hash"]):
        raise HTTPException(401, "E-posta veya parola hatalı")
    token = create_token(user["id"], user["role"])
    response.set_cookie("access_token", token, httponly=True, samesite="lax", max_age=ACCESS_TTL_MIN * 60, path="/")
    user.pop("_id", None); user.pop("password_hash", None)
    await audit(user, "login", "user", user["id"])
    return {"user": user, "access_token": token}


@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


@api.post("/auth/refresh")
async def refresh_session(user: dict = Depends(get_current_user), response: Response = None):
    token = create_token(user["id"], user["role"])
    if response is not None:
        response.set_cookie("access_token", token, httponly=True, samesite="lax", max_age=ACCESS_TTL_MIN * 60, path="/")
    return {"access_token": token, "user": user, "expires_in": ACCESS_TTL_MIN * 60}


@api.post("/auth/change-password")
async def change_password(inp: ChangePasswordIn, user: dict = Depends(get_current_user)):
    doc = await db.users.find_one({"id": user["id"]})
    if not doc or not verify_pw(inp.current_password, doc["password_hash"]):
        raise HTTPException(400, "Mevcut parola hatalı")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password_hash": hash_pw(inp.new_password)}})
    await audit(user, "password_change", "user", user["id"])
    return {"ok": True}


@api.post("/auth/forgot-password")
async def forgot_password(inp: ForgotPasswordIn):
    """Generate a reset token; send email. Always returns ok to avoid user enumeration."""
    doc = await db.users.find_one({"email": inp.email.lower()})
    if doc and doc.get("active", True):
        token = uuid.uuid4().hex + uuid.uuid4().hex  # 64 chars
        expires = datetime.now(timezone.utc) + timedelta(hours=1)
        await db.password_resets.insert_one({
            "id": new_id(),
            "user_id": doc["id"],
            "token": token,
            "expires_at": expires.isoformat(),
            "used": False,
            "created_at": now_iso(),
        })
        reset_url = f"{FRONTEND_URL}/reset-password?token={token}" if FRONTEND_URL else f"/reset-password?token={token}"
        html = (
            f'<table role="presentation" width="100%"><tr><td style="padding:24px;'
            f'font-family:Arial,sans-serif;color:#0f172a">'
            f'<h2 style="margin:0 0 12px 0">Parola Sıfırlama</h2>'
            f'<p>Merhaba {escape(doc.get("name") or "")},</p>'
            f'<p>{escape(EMAIL_FROM_NAME)} hesabınız için parola sıfırlama talebi aldık. '
            f'Aşağıdaki bağlantıya tıklayarak yeni parolanızı belirleyebilirsiniz. '
            f'Bağlantı 1 saat geçerlidir.</p>'
            f'<p style="margin:24px 0"><a href="{escape(reset_url)}" '
            f'style="background:#0f172a;color:#ffffff;padding:12px 20px;text-decoration:none;'
            f'border-radius:6px;display:inline-block">Parolamı Sıfırla</a></p>'
            f'<p style="font-size:12px;color:#64748b">Bu talebi siz yapmadıysanız bu e-postayı '
            f'yok sayabilirsiniz. {escape(EMAIL_FROM_NAME)} sizden e-posta ile asla parolanızı '
            f'veya kart bilgilerinizi istemez.</p>'
            f'<p style="font-size:12px;color:#94a3b8;margin-top:24px">— {escape(EMAIL_FROM_NAME)}</p>'
            f'</td></tr></table>'
        )
        try:
            await send_email(to=doc["email"], subject=f"{EMAIL_FROM_NAME} - Parola Sıfırlama", html=html)
        except Exception as e:
            log.error(f"Password reset email failed: {e}")
    return {"ok": True, "message": "Eğer e-posta sistemde kayıtlıysa sıfırlama bağlantısı gönderildi"}


@api.post("/auth/reset-password")
async def reset_password(inp: ResetPasswordIn):
    rec = await db.password_resets.find_one({"token": inp.token})
    if not rec:
        raise HTTPException(400, "Geçersiz sıfırlama bağlantısı")
    if rec.get("used"):
        raise HTTPException(400, "Bu bağlantı zaten kullanılmış")
    try:
        exp = datetime.fromisoformat(rec["expires_at"])
        if exp < datetime.now(timezone.utc):
            raise HTTPException(400, "Sıfırlama bağlantısının süresi dolmuş")
    except ValueError:
        raise HTTPException(400, "Geçersiz sıfırlama bağlantısı")
    await db.users.update_one({"id": rec["user_id"]}, {"$set": {"password_hash": hash_pw(inp.new_password)}})
    await db.password_resets.update_one({"id": rec["id"]}, {"$set": {"used": True, "used_at": now_iso()}})
    return {"ok": True}


@api.post("/auth/register")
async def register_user(inp: RegisterUserIn, admin: dict = Depends(require_role("admin"))):
    existing = await db.users.find_one({"email": inp.email.lower()})
    if existing:
        raise HTTPException(400, "Bu e-posta zaten kullanımda")
    user = {
        "id": new_id(),
        "email": inp.email.lower(),
        "name": inp.name,
        "role": inp.role,
        "active": True,
        "phone": inp.phone or "",
        "password_hash": hash_pw(inp.password),
        "created_at": now_iso(),
    }
    await db.users.insert_one(user)
    user.pop("_id", None); user.pop("password_hash", None)
    await audit(admin, "user_create", "user", user["id"], {"role": inp.role})
    return user


@api.get("/users")
async def list_users(admin: dict = Depends(require_role("admin"))):
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)
    return users


@api.patch("/users/{uid}")
async def update_user(uid: str, inp: UpdateUserIn, admin: dict = Depends(require_role("admin"))):
    updates = {}
    if inp.name is not None: updates["name"] = inp.name
    if inp.role is not None: updates["role"] = inp.role
    if inp.active is not None: updates["active"] = inp.active
    if inp.phone is not None: updates["phone"] = inp.phone
    if inp.password: updates["password_hash"] = hash_pw(inp.password)
    if not updates:
        raise HTTPException(400, "Güncellenecek alan yok")
    await db.users.update_one({"id": uid}, {"$set": updates})
    await audit(admin, "user_update", "user", uid, {"fields": list(updates.keys())})
    u = await db.users.find_one({"id": uid}, {"_id": 0, "password_hash": 0})
    return u


# ---------- Sites / Blocks / Units ----------
@api.get("/sites")
async def list_sites(user: dict = Depends(get_current_user)):
    return await db.sites.find({}, {"_id": 0}).to_list(1000)


@api.post("/sites")
async def create_site(inp: SiteIn, admin: dict = Depends(require_role("admin"))):
    site = {"id": new_id(), "created_at": now_iso(), **inp.model_dump()}
    await db.sites.insert_one(site)
    site.pop("_id", None)
    await audit(admin, "site_create", "site", site["id"], site_id=site["id"])
    return site


@api.patch("/sites/{sid}")
async def update_site(sid: str, inp: SiteIn, admin: dict = Depends(require_role("admin"))):
    await db.sites.update_one({"id": sid}, {"$set": inp.model_dump()})
    await audit(admin, "site_update", "site", sid, site_id=sid)
    return await db.sites.find_one({"id": sid}, {"_id": 0})


@api.get("/blocks")
async def list_blocks(site_id: str, user: dict = Depends(get_current_user)):
    return await db.blocks.find({"site_id": site_id}, {"_id": 0}).to_list(1000)


@api.post("/blocks")
async def create_block(inp: BlockIn, admin: dict = Depends(require_role("admin", "muhasebe"))):
    block = {"id": new_id(), "created_at": now_iso(), **inp.model_dump()}
    await db.blocks.insert_one(block)
    block.pop("_id", None)
    await audit(admin, "block_create", "block", block["id"], site_id=inp.site_id)
    return block


@api.delete("/blocks/{bid}")
async def delete_block(bid: str, admin: dict = Depends(require_role("admin"))):
    b = await db.blocks.find_one({"id": bid})
    if not b: raise HTTPException(404, "Blok bulunamadı")
    if await db.units.count_documents({"block_id": bid}) > 0:
        raise HTTPException(400, "Blokta bağımsız bölüm var, silinemez")
    await db.blocks.delete_one({"id": bid})
    await audit(admin, "block_delete", "block", bid, site_id=b["site_id"])
    return {"ok": True}


@api.get("/units")
async def list_units(site_id: str, user: dict = Depends(get_current_user)):
    units = await db.units.find({"site_id": site_id}, {"_id": 0}).to_list(2000)
    return units


@api.post("/units")
async def create_unit(inp: UnitIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    unit = {"id": new_id(), "created_at": now_iso(), **inp.model_dump()}
    await db.units.insert_one(unit)
    unit.pop("_id", None)
    await audit(u, "unit_create", "unit", unit["id"], site_id=inp.site_id)
    return unit


@api.patch("/units/{uid}")
async def update_unit(uid: str, inp: UnitPatchIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    unit = await db.units.find_one({"id": uid})
    if not unit: raise HTTPException(404, "Bağımsız bölüm bulunamadı")
    updates = {k: v for k, v in inp.model_dump().items() if v is not None}
    if not updates: raise HTTPException(400, "Güncellenecek alan yok")
    await db.units.update_one({"id": uid}, {"$set": updates})
    await audit(u, "unit_update", "unit", uid, {"fields": list(updates.keys())}, site_id=unit["site_id"])
    return await db.units.find_one({"id": uid}, {"_id": 0})


@api.delete("/units/{uid}")
async def delete_unit(uid: str, admin: dict = Depends(require_role("admin"))):
    unit = await db.units.find_one({"id": uid})
    if not unit: raise HTTPException(404, "Bağımsız bölüm bulunamadı")
    if await db.accruals.count_documents({"unit_id": uid}) > 0:
        raise HTTPException(400, "Bu birimde tahakkuk kayıtları var, silinemez")
    await db.units.delete_one({"id": uid})
    await audit(admin, "unit_delete", "unit", uid, site_id=unit["site_id"])
    return {"ok": True}


# ---------- Persons ----------
@api.get("/persons")
async def list_persons(site_id: str, user: dict = Depends(get_current_user)):
    return await db.persons.find({"site_id": site_id}, {"_id": 0}).to_list(2000)


@api.post("/persons")
async def create_person(inp: PersonIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    p = {"id": new_id(), "created_at": now_iso(), **inp.model_dump()}
    await db.persons.insert_one(p)
    p.pop("_id", None)
    await audit(u, "person_create", "person", p["id"], site_id=inp.site_id)
    return p


@api.patch("/persons/{pid}")
async def update_person(pid: str, inp: PersonPatchIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    person = await db.persons.find_one({"id": pid})
    if not person: raise HTTPException(404, "Kişi bulunamadı")
    updates = {k: v for k, v in inp.model_dump().items() if v is not None}
    if not updates: raise HTTPException(400, "Güncellenecek alan yok")
    await db.persons.update_one({"id": pid}, {"$set": updates})
    await audit(u, "person_update", "person", pid, {"fields": list(updates.keys())}, site_id=person["site_id"])
    return await db.persons.find_one({"id": pid}, {"_id": 0})


@api.get("/relations")
async def list_relations(site_id: str, unit_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {"site_id": site_id}
    if unit_id: q["unit_id"] = unit_id
    return await db.unit_person_relations.find(q, {"_id": 0}).to_list(2000)


@api.post("/relations")
async def create_relation(inp: RelationIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    r = {"id": new_id(), "created_at": now_iso(), **inp.model_dump()}
    await db.unit_person_relations.insert_one(r)
    r.pop("_id", None)
    await audit(u, "relation_create", "relation", r["id"], site_id=inp.site_id)
    return r


# ---------- Accounts (Cash + Bank) ----------
@api.get("/accounts")
async def list_accounts(site_id: str, user: dict = Depends(get_current_user)):
    cash = await db.cash_accounts.find({"site_id": site_id}, {"_id": 0}).to_list(200)
    bank = await db.bank_accounts.find({"site_id": site_id}, {"_id": 0}).to_list(200)
    for c in cash: c["kind"] = "cash"
    for b in bank: b["kind"] = "bank"
    # compute balances
    result = []
    for acc in cash + bank:
        balance = acc.get("opening_balance_kurus", 0)
        agg = db.account_transactions.aggregate([
            {"$match": {"site_id": site_id, "account_id": acc["id"], "account_kind": acc["kind"]}},
            {"$group": {"_id": None, "sum": {"$sum": "$signed_amount_kurus"}}},
        ])
        async for row in agg:
            balance += row["sum"]
        acc["balance_kurus"] = balance
        result.append(acc)
    return result


@api.post("/accounts")
async def create_account(inp: AccountIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    coll = db.cash_accounts if inp.kind == "cash" else db.bank_accounts
    acc = {"id": new_id(), "created_at": now_iso(), **inp.model_dump()}
    await coll.insert_one(acc)
    acc.pop("_id", None)
    await audit(u, "account_create", "account", acc["id"], {"kind": inp.kind}, site_id=inp.site_id)
    return acc


# ---------- Accruals (Tahakkuk) ----------
async def _get_advance(site_id: str, unit_id: str) -> int:
    doc = await db.advance_balances.find_one({"site_id": site_id, "unit_id": unit_id})
    return doc["balance_kurus"] if doc else 0


async def _set_advance(site_id: str, unit_id: str, new_balance: int):
    await db.advance_balances.update_one(
        {"site_id": site_id, "unit_id": unit_id},
        {"$set": {"balance_kurus": new_balance, "updated_at": now_iso()}},
        upsert=True,
    )


async def _write_journal(site_id: str, date_iso: str, description: str, lines: List[dict], user: dict, ref_type: str, ref_id: str):
    """lines: [{account_code, account_name, debit_kurus, credit_kurus, unit_id?}]"""
    total_debit = sum(l.get("debit_kurus", 0) for l in lines)
    total_credit = sum(l.get("credit_kurus", 0) for l in lines)
    if total_debit != total_credit:
        raise HTTPException(500, f"Yevmiye dengesizliği: borç {total_debit} ≠ alacak {total_credit}")
    entry = {
        "id": new_id(),
        "site_id": site_id,
        "date": date_iso,
        "description": description,
        "total_kurus": total_debit,
        "ref_type": ref_type,
        "ref_id": ref_id,
        "created_by": user["id"],
        "created_at": now_iso(),
        "lines": lines,
        "reversed": False,
        "reverses_id": None,
    }
    await db.journal_entries.insert_one(entry)
    entry.pop("_id", None)
    return entry


@api.post("/accruals/batch")
async def batch_accruals(inp: AccrualBatchIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    await check_period_open(inp.site_id, inp.due_date)
    units = await db.units.find({"site_id": inp.site_id}, {"_id": 0}).to_list(5000)
    if not units:
        raise HTTPException(400, "Bu sitede bağımsız bölüm yok")

    # calculate distribution
    per_unit: Dict[str, int] = {}
    if inp.method == "esit":
        for un in units:
            per_unit[un["id"]] = inp.amount_total_kurus
    else:
        key = "area_m2" if inp.method == "metrekare" else "share_ratio"
        total_weight = sum(un.get(key, 0) or 0 for un in units)
        if total_weight <= 0:
            raise HTTPException(400, f"Dağıtım için toplam {key} sıfır")
        allocated = 0
        for i, un in enumerate(units):
            w = un.get(key, 0) or 0
            share = int(round(inp.amount_total_kurus * w / total_weight))
            if i == len(units) - 1:
                share = inp.amount_total_kurus - allocated  # rounding correction
            per_unit[un["id"]] = share
            allocated += share

    created = []
    for un in units:
        amt = per_unit[un["id"]]
        if amt <= 0: continue
        acc = {
            "id": new_id(),
            "site_id": inp.site_id,
            "unit_id": un["id"],
            "period": inp.period,
            "description": inp.description,
            "due_date": inp.due_date,
            "amount_kurus": amt,
            "paid_kurus": 0,
            "status": "open",
            "reversed": False,
            "version": 1,
            "created_at": now_iso(),
        }
        await db.accruals.insert_one(acc)
        acc.pop("_id", None)
        # journal
        await _write_journal(
            inp.site_id, inp.due_date, f"Aidat Tahakkuku - {un['no']} - {inp.period}",
            [
                {"account_code": "120", "account_name": "Alacaklar/Cari", "debit_kurus": amt, "credit_kurus": 0, "unit_id": un["id"]},
                {"account_code": "600", "account_name": "Aidat Geliri", "debit_kurus": 0, "credit_kurus": amt},
            ],
            u, "accrual", acc["id"],
        )
        created.append(acc)

    await audit(u, "accrual_batch", "accrual", "", {"count": len(created), "period": inp.period}, site_id=inp.site_id)
    return {"created": len(created), "accruals": created}


@api.post("/accruals/extra")
async def extra_accrual(inp: AccrualExtraIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    await check_period_open(inp.site_id, inp.due_date)
    unit = await db.units.find_one({"id": inp.unit_id, "site_id": inp.site_id})
    if not unit: raise HTTPException(404, "Bağımsız bölüm bulunamadı")
    acc = {
        "id": new_id(), "site_id": inp.site_id, "unit_id": inp.unit_id,
        "period": inp.due_date[:7], "description": inp.description,
        "due_date": inp.due_date, "amount_kurus": inp.amount_kurus,
        "paid_kurus": 0, "status": "open", "reversed": False, "version": 1,
        "created_at": now_iso(),
    }
    await db.accruals.insert_one(acc)
    acc.pop("_id", None)
    await _write_journal(inp.site_id, inp.due_date, f"Ek Tahakkuk - {unit['no']} - {inp.description}",
        [
            {"account_code": "120", "account_name": "Alacaklar/Cari", "debit_kurus": inp.amount_kurus, "credit_kurus": 0, "unit_id": inp.unit_id},
            {"account_code": "600", "account_name": "Aidat Geliri", "debit_kurus": 0, "credit_kurus": inp.amount_kurus},
        ], u, "accrual", acc["id"])
    await audit(u, "accrual_extra", "accrual", acc["id"], site_id=inp.site_id)
    return acc


@api.get("/accruals")
async def list_accruals(site_id: str, unit_id: Optional[str] = None, status: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {"site_id": site_id}
    if unit_id: q["unit_id"] = unit_id
    if status: q["status"] = status
    return await db.accruals.find(q, {"_id": 0}).sort("due_date", 1).to_list(5000)


@api.post("/accruals/{aid}/reverse")
async def reverse_accrual(aid: str, inp: ReverseAccrualIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    acc = await db.accruals.find_one({"id": aid})
    if not acc: raise HTTPException(404, "Tahakkuk bulunamadı")
    if acc.get("reversed"): raise HTTPException(400, "Zaten iptal edilmiş")
    if acc.get("paid_kurus", 0) > 0: raise HTTPException(400, "Tahsilat yapılmış tahakkuk iptal edilemez")
    await check_period_open(acc["site_id"], acc["due_date"])
    await db.accruals.update_one({"id": aid}, {"$set": {"reversed": True, "status": "reversed"}})
    unit = await db.units.find_one({"id": acc["unit_id"]}) or {"no": ""}
    await _write_journal(acc["site_id"], now_iso()[:10], f"Tahakkuk İptali - {unit.get('no','')} - {inp.reason}",
        [
            {"account_code": "600", "account_name": "Aidat Geliri", "debit_kurus": acc["amount_kurus"], "credit_kurus": 0},
            {"account_code": "120", "account_name": "Alacaklar/Cari", "debit_kurus": 0, "credit_kurus": acc["amount_kurus"], "unit_id": acc["unit_id"]},
        ], u, "accrual_reverse", aid)
    await audit(u, "accrual_reverse", "accrual", aid, {"reason": inp.reason}, site_id=acc["site_id"])
    return {"ok": True}


# ---------- Collections (Tahsilat) with FIFO + Advance ----------
@api.post("/collections")
async def create_collection(inp: CollectionIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    await check_period_open(inp.site_id, inp.date)
    if inp.amount_kurus <= 0: raise HTTPException(400, "Tutar pozitif olmalı")

    # Check unique bank reference
    if inp.reference:
        exists = await db.account_transactions.find_one({
            "site_id": inp.site_id, "account_id": inp.account_id,
            "account_kind": inp.account_kind, "reference": inp.reference,
        })
        if exists:
            raise HTTPException(409, f"Mükerrer referans: '{inp.reference}' bu hesapta zaten var")

    unit = await db.units.find_one({"id": inp.unit_id, "site_id": inp.site_id})
    if not unit: raise HTTPException(404, "Bağımsız bölüm bulunamadı")

    # FIFO: apply to oldest unpaid accruals
    remaining = inp.amount_kurus
    open_accs = await db.accruals.find({
        "site_id": inp.site_id, "unit_id": inp.unit_id, "status": {"$in": ["open", "partial"]},
        "reversed": False,
    }, {"_id": 0}).sort("due_date", 1).to_list(1000)

    applications = []
    for acc in open_accs:
        if remaining <= 0: break
        unpaid = acc["amount_kurus"] - acc.get("paid_kurus", 0)
        take = min(unpaid, remaining)
        new_paid = acc.get("paid_kurus", 0) + take
        new_status = "paid" if new_paid >= acc["amount_kurus"] else "partial"
        # optimistic concurrency
        result = await db.accruals.update_one(
            {"id": acc["id"], "version": acc["version"]},
            {"$set": {"paid_kurus": new_paid, "status": new_status}, "$inc": {"version": 1}},
        )
        if result.matched_count == 0:
            raise HTTPException(409, "Bu kayıt başka bir kullanıcı tarafından güncellendi. Lütfen listeyi yenileyin.")
        applications.append({"accrual_id": acc["id"], "amount_kurus": take})
        remaining -= take

    # Excess -> advance
    advance_added = 0
    advance_used_note = ""
    if remaining > 0:
        cur = await _get_advance(inp.site_id, inp.unit_id)
        await _set_advance(inp.site_id, inp.unit_id, cur + remaining)
        advance_added = remaining
        advance_used_note = f"Fazla ödeme {remaining/100:.2f} TL avans hesabına aktarıldı"
        remaining = 0

    # Create payment record
    payment = {
        "id": new_id(), "site_id": inp.site_id, "unit_id": inp.unit_id,
        "account_id": inp.account_id, "account_kind": inp.account_kind,
        "date": inp.date, "amount_kurus": inp.amount_kurus,
        "reference": inp.reference or "", "note": inp.note or "",
        "applications": applications, "advance_added_kurus": advance_added,
        "created_by": u["id"], "created_at": now_iso(),
    }
    await db.payments.insert_one(payment)
    payment.pop("_id", None)

    # Account transaction
    await db.account_transactions.insert_one({
        "id": new_id(), "site_id": inp.site_id, "account_id": inp.account_id,
        "account_kind": inp.account_kind, "date": inp.date,
        "signed_amount_kurus": inp.amount_kurus,  # inflow
        "kind": "collection", "ref_id": payment["id"], "reference": inp.reference or "",
        "description": f"Tahsilat - {unit['no']}", "created_at": now_iso(),
    })

    # Journal: Debit Kasa/Banka, Credit Alacaklar (per applied accrual) + Credit Avans
    acc_name = "Kasa" if inp.account_kind == "cash" else "Bankalar"
    acc_code = "100" if inp.account_kind == "cash" else "102"
    lines = [{"account_code": acc_code, "account_name": acc_name, "debit_kurus": inp.amount_kurus, "credit_kurus": 0}]
    for ap in applications:
        lines.append({"account_code": "120", "account_name": "Alacaklar/Cari", "debit_kurus": 0, "credit_kurus": ap["amount_kurus"], "unit_id": inp.unit_id})
    if advance_added > 0:
        lines.append({"account_code": "340", "account_name": "Avanslar", "debit_kurus": 0, "credit_kurus": advance_added, "unit_id": inp.unit_id})
    await _write_journal(inp.site_id, inp.date, f"Tahsilat - {unit['no']}", lines, u, "payment", payment["id"])

    await audit(u, "collection_create", "payment", payment["id"], {"amount": inp.amount_kurus, "advance": advance_added}, site_id=inp.site_id)
    return {"payment": payment, "advance_added_kurus": advance_added, "message": advance_used_note or None}


@api.get("/collections")
async def list_collections(site_id: str, unit_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {"site_id": site_id}
    if unit_id: q["unit_id"] = unit_id
    return await db.payments.find(q, {"_id": 0}).sort("date", -1).to_list(2000)


# ---------- Excel Import: Collections (bank statement) ----------
from fastapi import UploadFile, File, Form

async def _parse_excel_rows(file: UploadFile) -> List[Dict[str, Any]]:
    """Parse uploaded xlsx to list of dicts using first row as headers (Turkish or English)."""
    from openpyxl import load_workbook
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    result = []
    for r in rows[1:]:
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in r):
            continue
        result.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
    return result


def _to_kurus(v) -> int:
    if v is None: return 0
    if isinstance(v, (int, float)):
        return int(round(float(v) * 100))
    s = str(v).strip().replace("₺", "").replace("TL", "").replace(" ", "")
    # Turkish: 1.234,56 or 1234,56 or 1234.56
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try: return int(round(float(s) * 100))
    except ValueError: return 0


def _parse_date(v) -> str:
    if v is None: return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    s = str(v).strip()
    # try YYYY-MM-DD
    for sep in ("-", "/", "."):
        parts = s.split(sep)
        if len(parts) == 3:
            if len(parts[0]) == 4:
                return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
            else:
                return f"{int(parts[2]):04d}-{int(parts[1]):02d}-{int(parts[0]):02d}"
    return s


@api.post("/collections/import")
async def import_collections(
    site_id: str = Form(...),
    account_id: str = Form(...),
    account_kind: str = Form(...),
    file: UploadFile = File(...),
    u: dict = Depends(require_role("admin", "muhasebe")),
):
    """
    Excel columns (headers required, Turkish): tarih, daire, tutar, referans, aciklama.
    Also accepts EN: date, unit, amount, reference, description.
    """
    if account_kind not in ("cash", "bank"):
        raise HTTPException(400, "Geçersiz hesap türü")
    rows = await _parse_excel_rows(file)
    if not rows:
        raise HTTPException(400, "Excel dosyası boş veya okunamadı")

    units = await db.units.find({"site_id": site_id}, {"_id": 0}).to_list(5000)
    unit_by_no = {str(un["no"]).strip(): un for un in units}

    def pick(row, keys):
        for k in keys:
            if k in row and row[k] not in (None, ""):
                return row[k]
        return None

    created = 0
    errors: List[dict] = []
    for idx, row in enumerate(rows, start=2):
        try:
            date_val = _parse_date(pick(row, ["tarih", "date"]))
            unit_no = str(pick(row, ["daire", "daire no", "bağımsız bölüm", "unit"]) or "").strip()
            amount_kurus = _to_kurus(pick(row, ["tutar", "amount"]))
            ref = str(pick(row, ["referans", "reference"]) or "").strip()
            note = str(pick(row, ["aciklama", "açıklama", "description", "not"]) or "").strip()
            if not date_val or not unit_no or amount_kurus <= 0:
                errors.append({"row": idx, "error": "Eksik alan: tarih/daire/tutar"}); continue
            unit = unit_by_no.get(unit_no)
            if not unit:
                errors.append({"row": idx, "error": f"Daire bulunamadı: {unit_no}"}); continue
            inp = CollectionIn(
                site_id=site_id, unit_id=unit["id"], account_id=account_id,
                account_kind=account_kind, date=date_val, amount_kurus=amount_kurus,
                reference=ref, note=note,
            )
            await create_collection(inp, u)
            created += 1
        except HTTPException as e:
            errors.append({"row": idx, "error": str(e.detail)})
        except Exception as e:
            errors.append({"row": idx, "error": str(e)})

    await audit(u, "collections_import", "payment", "", {"created": created, "errors": len(errors)}, site_id=site_id)
    return {"created": created, "errors": errors, "total": len(rows)}


@api.post("/accruals/import/natural-gas")
async def import_natural_gas(
    site_id: str = Form(...),
    period: str = Form(...),  # YYYY-MM
    due_date: str = Form(...),
    file: UploadFile = File(...),
    u: dict = Depends(require_role("admin", "muhasebe")),
):
    """
    Excel columns: daire, tuketim, tutar, aciklama (optional).
    Creates an 'extra' accrual for each row (description: 'Doğalgaz - <tuketim>').
    """
    await check_period_open(site_id, due_date)
    rows = await _parse_excel_rows(file)
    if not rows:
        raise HTTPException(400, "Excel dosyası boş veya okunamadı")

    units = await db.units.find({"site_id": site_id}, {"_id": 0}).to_list(5000)
    unit_by_no = {str(un["no"]).strip(): un for un in units}

    def pick(row, keys):
        for k in keys:
            if k in row and row[k] not in (None, ""):
                return row[k]
        return None

    created = 0
    errors: List[dict] = []
    for idx, row in enumerate(rows, start=2):
        try:
            unit_no = str(pick(row, ["daire", "daire no", "bağımsız bölüm", "unit"]) or "").strip()
            consumption = pick(row, ["tuketim", "tüketim", "consumption", "m3", "kwh"])
            amount_kurus = _to_kurus(pick(row, ["tutar", "amount", "bedel"]))
            note = str(pick(row, ["aciklama", "açıklama", "description", "not"]) or "").strip()
            if not unit_no or amount_kurus <= 0:
                errors.append({"row": idx, "error": "Eksik alan: daire/tutar"}); continue
            unit = unit_by_no.get(unit_no)
            if not unit:
                errors.append({"row": idx, "error": f"Daire bulunamadı: {unit_no}"}); continue
            desc_parts = ["Doğalgaz"]
            if consumption is not None: desc_parts.append(f"{consumption} m³")
            if note: desc_parts.append(note)
            inp = AccrualExtraIn(
                site_id=site_id, unit_id=unit["id"], due_date=due_date,
                description=" - ".join(desc_parts), amount_kurus=amount_kurus,
            )
            await extra_accrual(inp, u)
            created += 1
        except HTTPException as e:
            errors.append({"row": idx, "error": str(e.detail)})
        except Exception as e:
            errors.append({"row": idx, "error": str(e)})

    await audit(u, "natural_gas_import", "accrual", "", {"created": created, "period": period, "errors": len(errors)}, site_id=site_id)
    return {"created": created, "errors": errors, "total": len(rows)}


@api.get("/advances")
async def list_advances(site_id: str, user: dict = Depends(get_current_user)):
    return await db.advance_balances.find({"site_id": site_id, "balance_kurus": {"$gt": 0}}, {"_id": 0}).to_list(2000)


# ---------- Expenses ----------
@api.post("/expenses")
async def create_expense(inp: ExpenseIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    await check_period_open(inp.site_id, inp.date)
    if inp.reference:
        exists = await db.account_transactions.find_one({
            "site_id": inp.site_id, "account_id": inp.account_id,
            "account_kind": inp.account_kind, "reference": inp.reference,
        })
        if exists:
            raise HTTPException(409, f"Mükerrer referans: '{inp.reference}' bu hesapta zaten var")
    exp = {"id": new_id(), "created_by": u["id"], "created_at": now_iso(), **inp.model_dump()}
    await db.expenses.insert_one(exp)
    exp.pop("_id", None)
    await db.account_transactions.insert_one({
        "id": new_id(), "site_id": inp.site_id, "account_id": inp.account_id,
        "account_kind": inp.account_kind, "date": inp.date,
        "signed_amount_kurus": -inp.amount_kurus,
        "kind": "expense", "ref_id": exp["id"], "reference": inp.reference or "",
        "description": f"Gider - {inp.category}", "created_at": now_iso(),
    })
    acc_name = "Kasa" if inp.account_kind == "cash" else "Bankalar"
    acc_code = "100" if inp.account_kind == "cash" else "102"
    await _write_journal(inp.site_id, inp.date, f"Gider - {inp.category} - {inp.vendor or ''}",
        [
            {"account_code": "700", "account_name": f"Giderler ({inp.category})", "debit_kurus": inp.amount_kurus, "credit_kurus": 0},
            {"account_code": acc_code, "account_name": acc_name, "debit_kurus": 0, "credit_kurus": inp.amount_kurus},
        ], u, "expense", exp["id"])
    await audit(u, "expense_create", "expense", exp["id"], {"amount": inp.amount_kurus, "category": inp.category}, site_id=inp.site_id)
    return exp


@api.get("/expenses")
async def list_expenses(site_id: str, user: dict = Depends(get_current_user)):
    return await db.expenses.find({"site_id": site_id}, {"_id": 0}).sort("date", -1).to_list(5000)


# ---------- Transfers (Virman) ----------
@api.post("/transfers")
async def create_transfer(inp: TransferIn, u: dict = Depends(require_role("admin", "muhasebe"))):
    await check_period_open(inp.site_id, inp.date)
    if inp.amount_kurus <= 0: raise HTTPException(400, "Tutar pozitif olmalı")
    if inp.from_id == inp.to_id and inp.from_kind == inp.to_kind:
        raise HTTPException(400, "Kaynak ve hedef aynı olamaz")
    tid = new_id()
    await db.account_transactions.insert_one({
        "id": new_id(), "site_id": inp.site_id, "account_id": inp.from_id,
        "account_kind": inp.from_kind, "date": inp.date,
        "signed_amount_kurus": -inp.amount_kurus, "kind": "transfer_out",
        "ref_id": tid, "reference": inp.reference or "", "description": "Virman (çıkış)", "created_at": now_iso(),
    })
    await db.account_transactions.insert_one({
        "id": new_id(), "site_id": inp.site_id, "account_id": inp.to_id,
        "account_kind": inp.to_kind, "date": inp.date,
        "signed_amount_kurus": inp.amount_kurus, "kind": "transfer_in",
        "ref_id": tid, "reference": "", "description": "Virman (giriş)", "created_at": now_iso(),
    })
    from_code = "100" if inp.from_kind == "cash" else "102"
    to_code = "100" if inp.to_kind == "cash" else "102"
    from_name = "Kasa" if inp.from_kind == "cash" else "Bankalar"
    to_name = "Kasa" if inp.to_kind == "cash" else "Bankalar"
    await _write_journal(inp.site_id, inp.date, f"Virman - {inp.note or ''}",
        [
            {"account_code": to_code, "account_name": to_name, "debit_kurus": inp.amount_kurus, "credit_kurus": 0},
            {"account_code": from_code, "account_name": from_name, "debit_kurus": 0, "credit_kurus": inp.amount_kurus},
        ], u, "transfer", tid)
    await audit(u, "transfer_create", "transfer", tid, {"amount": inp.amount_kurus}, site_id=inp.site_id)
    return {"id": tid}


@api.get("/account-transactions")
async def list_account_transactions(site_id: str, account_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {"site_id": site_id}
    if account_id: q["account_id"] = account_id
    return await db.account_transactions.find(q, {"_id": 0}).sort("date", -1).to_list(5000)


# ---------- Journal ----------
@api.get("/journal")
async def list_journal(site_id: str, limit: int = 200, user: dict = Depends(get_current_user)):
    return await db.journal_entries.find({"site_id": site_id}, {"_id": 0}).sort("date", -1).limit(limit).to_list(limit)


# ---------- Periods ----------
@api.get("/periods")
async def list_periods(site_id: str, user: dict = Depends(get_current_user)):
    return await db.periods.find({"site_id": site_id}, {"_id": 0}).sort([("year", -1), ("month", -1)]).to_list(500)


@api.post("/periods/close")
async def close_period(inp: PeriodActionIn, admin: dict = Depends(require_role("admin"))):
    await db.periods.update_one(
        {"site_id": inp.site_id, "year": inp.year, "month": inp.month},
        {"$set": {"status": "closed", "closed_at": now_iso(), "closed_by": admin["id"]}},
        upsert=True,
    )
    await audit(admin, "period_close", "period", f"{inp.year}-{inp.month:02d}", {"reason": inp.reason}, site_id=inp.site_id)
    return {"ok": True}


@api.post("/periods/reopen")
async def reopen_period(inp: PeriodActionIn, admin: dict = Depends(require_role("admin"))):
    if not inp.reason or len(inp.reason.strip()) < 5:
        raise HTTPException(400, "Yeniden açma için gerekçe zorunludur (en az 5 karakter)")
    await db.periods.update_one(
        {"site_id": inp.site_id, "year": inp.year, "month": inp.month},
        {"$set": {"status": "open", "reopened_at": now_iso(), "reopened_by": admin["id"], "reopen_reason": inp.reason}},
        upsert=True,
    )
    await audit(admin, "period_reopen", "period", f"{inp.year}-{inp.month:02d}", {"reason": inp.reason}, site_id=inp.site_id)
    return {"ok": True}


# ---------- Audit ----------
@api.get("/audit")
async def list_audit(site_id: Optional[str] = None, limit: int = 200, user: dict = Depends(get_current_user)):
    q = {}
    if site_id: q["site_id"] = site_id
    return await db.audit_logs.find(q, {"_id": 0}).sort("ts", -1).limit(limit).to_list(limit)


# ---------- Dashboard ----------
@api.get("/dashboard")
async def dashboard(site_id: str, user: dict = Depends(get_current_user)):
    # totals
    unit_count = await db.units.count_documents({"site_id": site_id})
    person_count = await db.persons.count_documents({"site_id": site_id, "active": True})

    # total outstanding debt
    agg = db.accruals.aggregate([
        {"$match": {"site_id": site_id, "reversed": False, "status": {"$in": ["open", "partial"]}}},
        {"$group": {"_id": None, "debt": {"$sum": {"$subtract": ["$amount_kurus", "$paid_kurus"]}}}},
    ])
    debt = 0
    async for r in agg: debt = r["debt"]

    # this month collections/expenses
    today = datetime.now(timezone.utc)
    period_prefix = f"{today.year:04d}-{today.month:02d}"
    coll_agg = db.payments.aggregate([
        {"$match": {"site_id": site_id, "date": {"$regex": f"^{period_prefix}"}}},
        {"$group": {"_id": None, "sum": {"$sum": "$amount_kurus"}}},
    ])
    coll_sum = 0
    async for r in coll_agg: coll_sum = r["sum"]

    exp_agg = db.expenses.aggregate([
        {"$match": {"site_id": site_id, "date": {"$regex": f"^{period_prefix}"}}},
        {"$group": {"_id": None, "sum": {"$sum": "$amount_kurus"}}},
    ])
    exp_sum = 0
    async for r in exp_agg: exp_sum = r["sum"]

    # cash+bank total
    accounts = await list_accounts(site_id, user)
    cash_bank_total = sum(a["balance_kurus"] for a in accounts)

    # debtor count
    debtor_agg = db.accruals.aggregate([
        {"$match": {"site_id": site_id, "reversed": False, "status": {"$in": ["open", "partial"]}}},
        {"$group": {"_id": "$unit_id"}},
        {"$count": "count"},
    ])
    debtor_count = 0
    async for r in debtor_agg: debtor_count = r["count"]

    return {
        "unit_count": unit_count,
        "person_count": person_count,
        "total_debt_kurus": debt,
        "month_collections_kurus": coll_sum,
        "month_expenses_kurus": exp_sum,
        "cash_bank_total_kurus": cash_bank_total,
        "debtor_count": debtor_count,
    }


@api.get("/dashboard/activity")
async def dashboard_activity(site_id: str, days: int = 7, user: dict = Depends(get_current_user)):
    """Returns last N days of daily collections and expenses."""
    days = max(1, min(days, 90))
    today = datetime.now(timezone.utc).date()
    start = today - timedelta(days=days - 1)
    days_list = [(start + timedelta(days=i)).isoformat() for i in range(days)]
    coll_map = {d: 0 for d in days_list}
    exp_map = {d: 0 for d in days_list}

    async for row in db.payments.aggregate([
        {"$match": {"site_id": site_id, "date": {"$gte": days_list[0], "$lte": days_list[-1]}}},
        {"$group": {"_id": {"$substr": ["$date", 0, 10]}, "sum": {"$sum": "$amount_kurus"}}},
    ]):
        if row["_id"] in coll_map:
            coll_map[row["_id"]] = row["sum"]

    async for row in db.expenses.aggregate([
        {"$match": {"site_id": site_id, "date": {"$gte": days_list[0], "$lte": days_list[-1]}}},
        {"$group": {"_id": {"$substr": ["$date", 0, 10]}, "sum": {"$sum": "$amount_kurus"}}},
    ]):
        if row["_id"] in exp_map:
            exp_map[row["_id"]] = row["sum"]

    return [
        {"date": d, "collections_kurus": coll_map[d], "expenses_kurus": exp_map[d]}
        for d in days_list
    ]


# ---------- Reports ----------
async def _debtor_list(site_id: str):
    pipeline = [
        {"$match": {"site_id": site_id, "reversed": False, "status": {"$in": ["open", "partial"]}}},
        {"$group": {"_id": "$unit_id", "debt": {"$sum": {"$subtract": ["$amount_kurus", "$paid_kurus"]}}, "count": {"$sum": 1}}},
        {"$sort": {"debt": -1}},
    ]
    rows = []
    async for row in db.accruals.aggregate(pipeline):
        unit = await db.units.find_one({"id": row["_id"]}, {"_id": 0}) or {}
        rows.append({"unit_id": row["_id"], "unit_no": unit.get("no", ""), "debt_kurus": row["debt"], "count": row["count"]})
    return rows


@api.get("/reports/debtors")
async def report_debtors(site_id: str, user: dict = Depends(get_current_user)):
    return await _debtor_list(site_id)


@api.get("/reports/unit-statement")
async def unit_statement(site_id: str, unit_id: str, user: dict = Depends(get_current_user)):
    accruals = await db.accruals.find({"site_id": site_id, "unit_id": unit_id}, {"_id": 0}).sort("due_date", 1).to_list(2000)
    payments = await db.payments.find({"site_id": site_id, "unit_id": unit_id}, {"_id": 0}).sort("date", 1).to_list(2000)
    advance = await _get_advance(site_id, unit_id)
    return {"accruals": accruals, "payments": payments, "advance_kurus": advance}


@api.get("/reports/income-expense")
async def income_expense(site_id: str, period: Optional[str] = None, user: dict = Depends(get_current_user)):
    match: Dict[str, Any] = {"site_id": site_id}
    if period:
        match["date"] = {"$regex": f"^{period}"}
    inc_agg = db.payments.aggregate([{"$match": match}, {"$group": {"_id": None, "sum": {"$sum": "$amount_kurus"}}}])
    exp_agg = db.expenses.aggregate([{"$match": match}, {"$group": {"_id": "$category", "sum": {"$sum": "$amount_kurus"}}}])
    inc = 0
    async for r in inc_agg: inc = r["sum"]
    exp_by_cat = []
    async for r in exp_agg: exp_by_cat.append({"category": r["_id"], "sum_kurus": r["sum"]})
    exp_total = sum(e["sum_kurus"] for e in exp_by_cat)
    return {"income_kurus": inc, "expense_kurus": exp_total, "expenses_by_category": exp_by_cat, "net_kurus": inc - exp_total}


def _fmt_tl(kurus: int) -> str:
    return f"{kurus/100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " TL"


@api.get("/reports/debtors/export.xlsx")
async def debtors_xlsx(site_id: str, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    rows = await _debtor_list(site_id)
    wb = Workbook(); ws = wb.active; ws.title = "Borçlular"
    ws.append(["Bağımsız Bölüm No", "Borç Tutarı (TL)", "Açık Tahakkuk Sayısı"])
    for r in rows:
        ws.append([r["unit_no"], round(r["debt_kurus"]/100, 2), r["count"]])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="borclular.xlsx"'})


@api.get("/reports/debtors/export.pdf")
async def debtors_pdf(site_id: str, user: dict = Depends(get_current_user)):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm
    rows = await _debtor_list(site_id)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 14); c.drawString(2*cm, h - 2*cm, "Borçlu Listesi")
    c.setFont("Helvetica", 9); c.drawString(2*cm, h - 2.6*cm, f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    y = h - 3.5*cm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(2*cm, y, "B.B. No"); c.drawString(6*cm, y, "Borç"); c.drawString(11*cm, y, "Açık Tahakkuk")
    y -= 0.5*cm; c.line(2*cm, y, 19*cm, y); y -= 0.4*cm
    c.setFont("Helvetica", 9)
    total = 0
    for r in rows:
        if y < 2*cm:
            c.showPage(); y = h - 2*cm; c.setFont("Helvetica", 9)
        c.drawString(2*cm, y, str(r["unit_no"]))
        c.drawString(6*cm, y, _fmt_tl(r["debt_kurus"]))
        c.drawString(11*cm, y, str(r["count"]))
        total += r["debt_kurus"]
        y -= 0.5*cm
    y -= 0.3*cm; c.setFont("Helvetica-Bold", 10)
    c.drawString(2*cm, y, "TOPLAM"); c.drawString(6*cm, y, _fmt_tl(total))
    c.save(); buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": 'attachment; filename="borclular.pdf"'})


@api.get("/reports/income-expense/export.xlsx")
async def ie_xlsx(site_id: str, period: Optional[str] = None, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    data = await income_expense(site_id, period, user)
    wb = Workbook(); ws = wb.active; ws.title = "Gelir-Gider"
    ws.append(["Kategori", "Tutar (TL)"])
    ws.append(["Toplam Gelir (Tahsilat)", round(data["income_kurus"]/100, 2)])
    for e in data["expenses_by_category"]:
        ws.append([f"Gider: {e['category']}", round(e['sum_kurus']/100, 2)])
    ws.append(["Toplam Gider", round(data["expense_kurus"]/100, 2)])
    ws.append(["Net", round(data["net_kurus"]/100, 2)])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="gelir-gider.xlsx"'})


# ---------- Sync (offline write queue) ----------
class SyncBatchIn(BaseModel):
    operations: List[Dict[str, Any]]


@api.post("/sync/batch")
async def sync_batch(inp: SyncBatchIn, request: Request, user: dict = Depends(get_current_user)):
    """Process offline-queued operations. Each op: {client_id, endpoint, method, body}"""
    results = []
    from httpx import AsyncClient
    # process inline by calling internal handlers via HTTP loopback is heavy; instead, dispatch manually
    for op in inp.operations:
        cid = op.get("client_id", new_id())
        endpoint = op.get("endpoint", "")
        body = op.get("body", {})
        try:
            if endpoint == "/collections":
                res = await create_collection(CollectionIn(**body), user)
            elif endpoint == "/expenses":
                res = await create_expense(ExpenseIn(**body), user)
            elif endpoint == "/accruals/extra":
                res = await extra_accrual(AccrualExtraIn(**body), user)
            else:
                results.append({"client_id": cid, "ok": False, "error": f"Bilinmeyen endpoint: {endpoint}"})
                continue
            results.append({"client_id": cid, "ok": True, "data": res})
        except HTTPException as e:
            results.append({"client_id": cid, "ok": False, "error": e.detail, "status": e.status_code})
        except Exception as e:
            results.append({"client_id": cid, "ok": False, "error": str(e)})
    return {"results": results}


# ---------- Include ----------
app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)
